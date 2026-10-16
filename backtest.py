"""Evaluate Top30-Bottom30 portfolios using historical rolling four-factor regressions."""

from __future__ import annotations

import argparse
from datetime import datetime
import os
from pathlib import Path
import re
import sys
import tempfile

import numpy as np
import pandas as pd


ROOT = Path(__file__).resolve().parent
FACTORS = ["mktrf", "smb", "hml", "mom"]
METRICS = ["AR", "SR", "Sortino", "Annual abnormal return"]
TOP_N = 30


def resolve_path(path: Path) -> Path:
    return path.resolve() if path.is_absolute() else (ROOT / path).resolve()


def parse_months(values: pd.Series) -> pd.PeriodIndex:
    text = values.astype(str).str.strip()
    if not text.str.fullmatch(r"\d{6}").all():
        raise ValueError("Months must use the YYYYMM format.")
    dates = pd.to_datetime(text, format="%Y%m", errors="raise")
    return pd.PeriodIndex(dates, freq="M")


def load_predictions(path: Path) -> tuple[dict[str, pd.DataFrame], list[int]]:
    if path.is_dir():
        files = sorted(p for p in path.glob("*_test_predictions.csv")
                       if not p.name.startswith("._"))
    else:
        files = [path]
    if not files:
        raise ValueError(f"No test prediction files found: {path}")
    per_seed = {}
    expected_months = None
    required = {"date", "instrument", "y_pred", "y_true"}
    for file in files:
        frame = pd.read_csv(file, dtype={"instrument": str})
        if not required.issubset(frame.columns):
            raise ValueError(f"Missing prediction columns in {file.name}: {required - set(frame.columns)}")
        periods = parse_months(frame["date"])
        frame["date"] = periods.year * 100 + periods.month
        if frame["instrument"].isna().any():
            raise ValueError(f"Missing fund identifier in {file.name}")
        frame["instrument"] = frame["instrument"].str.strip()
        if frame["instrument"].eq("").any():
            raise ValueError(f"Empty fund identifier in {file.name}")
        for column in ("y_pred", "y_true"):
            frame[column] = pd.to_numeric(frame[column], errors="raise")
        match = re.search(r"_(?:seed|s)(\d+)_test_predictions$", file.stem)
        name = str(int(match.group(1))) if match else file.stem
        groups = frame.groupby("seed", sort=True) if "seed" in frame else [(name, frame)]
        for seed, group in groups:
            key = str(seed)
            if key in per_seed:
                raise ValueError(f"Duplicate seed: {key}")
            if group.duplicated(["date", "instrument"]).any():
                raise ValueError(f"Duplicate prediction keys for seed {key}")
            months = sorted(int(value) for value in group["date"].unique())
            if len(months) < 2:
                raise ValueError("At least two evaluation months are required.")
            periods = parse_months(pd.Series(months))
            if not np.array_equal(periods.asi8, pd.period_range(periods[0], periods[-1], freq="M").asi8):
                raise ValueError(f"Missing evaluation month for seed {key}")
            if expected_months is None:
                expected_months = months
            elif months != expected_months:
                raise ValueError("Evaluation months differ across seeds.")
            per_seed[key] = group[list(required)].copy()
    if not per_seed:
        raise ValueError("No prediction groups found.")
    order = sorted(per_seed, key=lambda seed: (0, int(seed)) if seed.isdigit() else (1, seed))
    return {seed: per_seed[seed] for seed in order}, expected_months


def rolling_funds(data_dir: Path, evaluation_months: list[int]) -> pd.DataFrame:
    """Use the original grouped least-squares calculation over [t-24, t-1]."""
    returns = pd.read_csv(data_dir / "fundreturn.csv", comment="#", dtype={"instrument": str})
    needed = {"date", "instrument", "fund_return_net"}
    if not needed.issubset(returns.columns):
        raise ValueError("fundreturn.csv requires date, instrument, and fund_return_net.")
    returns["month"] = parse_months(returns["date"])
    returns["fund_return_net"] = pd.to_numeric(returns["fund_return_net"], errors="raise")
    if returns["instrument"].isna().any():
        raise ValueError("Missing return-data fund identifier.")
    returns["instrument"] = returns["instrument"].str.strip()
    if returns["instrument"].eq("").any():
        raise ValueError("Empty return-data fund identifier.")
    if returns.duplicated(["month", "instrument"]).any():
        raise ValueError("Duplicate fund-month returns.")
    first = pd.Period(str(evaluation_months[0]), freq="M") - 24
    last = pd.Period(str(evaluation_months[-1]), freq="M")
    calendar = pd.period_range(first, last, freq="M")
    returns = returns.loc[returns["month"].between(first, last)]
    wide = returns.pivot(index="month", columns="instrument", values="fund_return_net").reindex(calendar)
    factors = pd.read_csv(data_dir / "factors_monthly.csv")
    if not {"date", "rf", *FACTORS}.issubset(factors.columns):
        raise ValueError("factors_monthly.csv requires date, rf, mktrf, smb, hml, and mom.")
    factors["month"] = parse_months(factors["date"])
    if factors["month"].duplicated().any():
        raise ValueError("Duplicate factor months.")
    factors = factors.set_index("month").reindex(calendar)
    x = factors[FACTORS].to_numpy(dtype=float)
    rf = factors["rf"].to_numpy(dtype=float)
    values = wide.to_numpy(dtype=float)
    if not np.isfinite(np.column_stack([x, rf])).all():
        raise ValueError(f"Complete finite factors and RF are required from {first} to {last}.")
    frames = []
    for date in evaluation_months:
        t = calendar.get_loc(pd.Period(str(date), freq="M"))
        history = values[t - 24:t]
        valid = np.isfinite(history)
        nobs = valid.sum(axis=0)
        eligible = np.flatnonzero((nobs >= 18) & np.isfinite(values[t]))
        design = np.column_stack([np.ones(24), x[t - 24:t]])
        patterns, membership = np.unique(valid[:, eligible].T, axis=0, return_inverse=True)
        for group, mask in enumerate(patterns):
            columns = eligible[membership == group]
            xx = design[mask]
            if np.linalg.matrix_rank(xx) != 5:
                continue
            yy = history[mask][:, columns] - rf[t - 24:t][mask, None]
            coefficient = np.linalg.lstsq(xx, yy, rcond=None)[0]
            # Preserve the independent numerical checks from the reference script.
            q, rr = np.linalg.qr(xx, mode="reduced")
            independent = np.linalg.solve(rr, q.T @ yy[:, 0])
            np.testing.assert_allclose(coefficient[:, 0], independent, atol=1e-10, rtol=1e-9)
            np.testing.assert_allclose((yy - xx @ coefficient).mean(axis=0), 0, atol=1e-11)
            abnormal = values[t, columns] - rf[t] - x[t] @ coefficient[1:]
            frames.append(pd.DataFrame({
                "date": date, "instrument": wide.columns[columns],
                "fund_return_net": values[t, columns], "n_history": nobs[columns],
                "abnormal": abnormal,
            }))
    if not frames:
        raise ValueError("No fund-months satisfy the historical regression conditions.")
    result = pd.concat(frames, ignore_index=True)
    if not np.isfinite(result[["fund_return_net", "abnormal"]].to_numpy()).all():
        raise ValueError("Nonfinite rolling regression results.")
    return result


def calculate_seed(prediction: pd.DataFrame, funds: pd.DataFrame,
                   months: list[int]) -> dict:
    """Retain the reference strategy and the four requested performance metrics."""
    prediction = prediction.loc[np.isfinite(prediction["y_pred"])
                                & np.isfinite(prediction["y_true"])].copy()
    matched = prediction.merge(funds, on=["date", "instrument"], validate="one_to_one")
    counts = matched.groupby("date").size().reindex(months, fill_value=0)
    if counts.min() < 2 * TOP_N:
        raise ValueError("Fewer than 60 eligible funds after label/history filtering.")
    ranked = matched.sort_values(["date", "y_pred", "instrument"],
                                 ascending=[True, False, True], kind="mergesort")
    groups = ranked.groupby("date", sort=True)
    top, bottom = groups.head(TOP_N), groups.tail(TOP_N)
    keys = ["date", "instrument"]
    expected_count = len(months) * TOP_N
    if len(top) != expected_count or len(bottom) != expected_count or len(top.merge(bottom, on=keys)):
        raise ValueError("Invalid or overlapping holdings.")
    columns = ["fund_return_net", "abnormal"]
    difference = top.groupby("date")[columns].mean() - bottom.groupby("date")[columns].mean()
    monthly, abnormal = difference.reindex(months).to_numpy().T
    signed = pd.concat([top.assign(weight=1 / TOP_N), bottom.assign(weight=-1 / TOP_N)])
    for column, expected in zip(columns, [monthly, abnormal]):
        actual = (signed[column] * signed["weight"]).groupby(signed["date"]).sum().reindex(months)
        np.testing.assert_allclose(actual, expected, atol=1e-12, rtol=1e-10)
    if not np.isfinite(difference.to_numpy()).all() or np.any(1 + monthly <= 0):
        raise ValueError("Invalid portfolio returns.")
    sd = np.std(monthly, ddof=1)
    downside = np.sqrt(np.mean(np.minimum(monthly, 0) ** 2))
    # Sharpe is undefined when monthly long-short returns have zero
    # standard deviation; leave the value blank.
    if sd == 0:
        sharpe = None
    else:
        sharpe = np.sqrt(12) * monthly.mean() / sd
    # Sortino is undefined when all monthly long-short returns
    # during the backtest period are nonnegative; leave the value blank.
    if downside == 0:
        sortino = None
    else:
        sortino = np.sqrt(12) * monthly.mean() / downside
    metrics = dict(zip(METRICS, [
        12 * monthly.mean(), sharpe, sortino, 12 * abnormal.mean(),
    ]))
    defined_metrics = [value for value in metrics.values() if value is not None]
    if not np.isfinite(defined_metrics).all():
        raise ValueError("Nonfinite performance metrics.")
    cumulative = np.cumprod(1 + monthly) - 1
    cumulative_abnormal = np.cumsum(abnormal)
    np.testing.assert_allclose(cumulative_abnormal[-1],
                               len(months) / 12 * metrics["Annual abnormal return"], atol=1e-12)
    return {"metrics": metrics, "cumulative_return": cumulative,
            "cumulative_abnormal_return": cumulative_abnormal}


def write_workbook(results: dict[str, dict], months: list[int], destination: Path) -> None:
    """Write the only backtest output: one workbook with four numeric worksheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append(["metric", "n_seeds", "mean", "sample_std"])
    for metric in METRICS:
        values = np.asarray([result["metrics"][metric] for result in results.values()
                             if result["metrics"][metric] is not None], dtype=float)
        summary.append([metric, len(values), float(values.mean()) if len(values) else None,
                        float(values.std(ddof=1)) if len(values) > 1 else None])
        number_format = "0.000" if metric in {"SR", "Sortino"} else "0.00%"
        for cell in summary[summary.max_row][2:]:
            cell.number_format = number_format
    per_seed = book.create_sheet("PerSeedMetrics")
    per_seed.append(["seed", "start_month", "end_month", "n_months", *METRICS])
    dates = [datetime.strptime(str(month), "%Y%m") for month in months]
    for seed, result in results.items():
        per_seed.append([seed, dates[0], dates[-1], len(months),
                         *[float(result["metrics"][name])
                           if result["metrics"][name] is not None else None
                           for name in METRICS]])
        row = per_seed.max_row
        for column in (2, 3):
            per_seed.cell(row, column).number_format = "yyyy-mm"
        for column, name in enumerate(METRICS, 5):
            per_seed.cell(row, column).number_format = "0.000" if name in {"SR", "Sortino"} else "0.00%"
    for title, field in [("CumulativeReturn", "cumulative_return"),
                         ("CumulativeAbnormalReturn", "cumulative_abnormal_return")]:
        sheet = book.create_sheet(title)
        sheet.append(["date", *[f"seed{seed}" for seed in results], "mean"])
        matrix = np.asarray([result[field] for result in results.values()])
        means = matrix.mean(axis=0)
        for index, date in enumerate(dates):
            sheet.append([date, *[float(value) for value in matrix[:, index]], float(means[index])])
            sheet.cell(sheet.max_row, 1).number_format = "yyyy-mm"
            for cell in sheet[sheet.max_row][1:]:
                cell.number_format = "0.00%"
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="203864")
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 24
        for column in sheet.columns:
            width = max(len(str(cell.value)) if cell.row == 1 or isinstance(cell.value, str)
                        else 12 for cell in column) + 3
            sheet.column_dimensions[column[0].column_letter].width = min(max(width, 14), 32)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Calibri", size=11)
    destination.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(prefix=".backtest_", suffix=".xlsx", dir=destination.parent)
    os.close(handle)
    temporary = Path(temporary_name)
    try:
        book.save(temporary)
        os.replace(temporary, destination)
    finally:
        book.close()
        if temporary.exists():
            temporary.unlink()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--predictions", type=Path, default=ROOT / "outputs/test_predictions")
    parser.add_argument("--data-dir", type=Path, default=ROOT / "data")
    parser.add_argument("--output", type=Path, default=ROOT / "outputs/backtest.xlsx")
    args = parser.parse_args(argv)
    destination = resolve_path(args.output)
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("The output must be one .xlsx workbook.")
    groups, months = load_predictions(resolve_path(args.predictions))
    funds = rolling_funds(resolve_path(args.data_dir), months)
    results = {}
    for seed, prediction in groups.items():
        results[seed] = calculate_seed(prediction, funds, months)
        print(f"Evaluated seed {seed}: {len(months)} months", flush=True)
    write_workbook(results, months, destination)
    print(f"backtest={destination}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
